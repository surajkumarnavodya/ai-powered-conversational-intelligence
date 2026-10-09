"""HTTP tests for `api/bi_integration.py` (Prompt 48,
`48_ENTERPRISE_BI_INTEGRATIONS_CONTRACT.md`) -- a real in-memory SQLite
identity database (real tenants, real local accounts, real JWT login),
mirroring `tests/test_api_eval_platform.py`/`tests/test_api_governance.py`'s
own fixture precedent, combined with `tests/test_api_execute.py`'s own
mocked-SQL-execution pattern (`execute_readonly_sql`/`get_read_only_engine`
patched at the `api.bi_integration` module they're looked up from) -- the
mission's own explicit "Test authorization and refresh failures"
requirement is covered by `TestRefreshFailures`/`TestUnauthorizedAccess`
below.
"""

from __future__ import annotations

from collections.abc import Iterator
from pathlib import Path

import identity.db as identity_db_mod
import pytest
from fastapi.testclient import TestClient
from identity.bootstrap import seed_rbac
from identity.models import Base
from identity.repositories.tenants import create_tenant
from identity.repositories.users import create_user
from sqlalchemy import create_engine
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.pool import StaticPool

import api.auth as api_auth_mod
import api.bi_integration as bi_integration_mod
import api.identity_auth as identity_auth_mod
import api.identity_authz as identity_authz_mod
import api.main as api_main
import api.rate_limit as api_rate_limit_mod
from bi_integration.push import PushResult
from config.settings import Settings
from security.secrets import SecretStr

PASSWORD = "correcthorse battery staple 1"

_SETTINGS = Settings(
    local_auth_enabled=True,
    auth_database_url="postgresql://placeholder/unused",
    jwt_secret_key="s" * 40,
    jwt_issuer="text-to-sql-agent",
    jwt_audience="text-to-sql-web",
    allow_public_registration=True,
    cookie_secure=False,
    password_min_length=8,
    login_rate_limit_per_minute=1000,
    register_rate_limit_per_hour=1000,
    api_action_rate_limit_per_minute=1000,
    enable_bi_integration=True,
    db_type="postgresql",
    db_host="db.example.com",
    db_port=5432,
    db_name="mydb",
    db_user="reader",
    db_password=SecretStr("secret"),
    db_connection_string=None,
    db_schema=None,
    db_odbc_driver="x",
    chroma_persist_dir=Path("/tmp/chroma"),
    chroma_collection_name="schema_ddl",
    embedding_model_name="all-MiniLM-L6-v2",
    query_timeout_seconds=15,
    bi_export_max_rows=1000,
)


@pytest.fixture(autouse=True)
def _identity_test_db(monkeypatch) -> Iterator[object]:
    engine = create_engine(
        "sqlite:///:memory:", connect_args={"check_same_thread": False}, poolclass=StaticPool
    )
    Base.metadata.create_all(engine)

    monkeypatch.setattr(identity_db_mod, "get_identity_engine", lambda settings=None: engine)
    monkeypatch.setattr(api_auth_mod, "get_settings", lambda: _SETTINGS)
    monkeypatch.setattr(identity_auth_mod, "get_settings", lambda: _SETTINGS)
    monkeypatch.setattr(identity_authz_mod, "get_settings", lambda: _SETTINGS)
    monkeypatch.setattr(bi_integration_mod, "get_settings", lambda: _SETTINGS)
    monkeypatch.setattr(bi_integration_mod, "get_read_only_engine", lambda config: object())

    session = identity_db_mod.get_identity_session(_SETTINGS)
    seed_rbac(session)
    create_tenant(session, tenant_id="tenant-a", name="Tenant A")
    create_tenant(session, tenant_id="tenant-b", name="Tenant B")
    session.commit()
    session.close()

    api_rate_limit_mod._limiters.clear()
    identity_auth_mod._login_limiters.clear()
    identity_auth_mod._register_limiters.clear()
    yield engine


@pytest.fixture
def client() -> TestClient:
    return TestClient(api_main.app)


def _account(client: TestClient, email: str, role: str, tenant_id: str = "tenant-a") -> dict:
    session = identity_db_mod.get_identity_session(_SETTINGS)
    try:
        create_user(session, email=email, password=PASSWORD, tenant_id=tenant_id, role_name=role)
        session.commit()
    finally:
        session.close()
    response = client.post("/auth/login", json={"email": email, "password": PASSWORD})
    assert response.status_code == 200, response.text
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


def _create_connector(client: TestClient, headers: dict, **overrides: object) -> dict:
    payload = {
        "target_type": "power_bi",
        "display_name": "Revenue Push",
        "sql_text": "SELECT region, revenue FROM sales",
        "push_config": {"dataset_id": "ds1", "table_name": "tbl1"},
    }
    payload.update(overrides)
    response = client.post("/bi/connectors", json=payload, headers=headers)
    assert response.status_code == 200, response.text
    return response.json()


def _mock_execute_success(monkeypatch, columns=("region", "revenue"), rows=(("East", 100),)):
    monkeypatch.setattr(
        bi_integration_mod,
        "execute_readonly_sql",
        lambda sql, timeout, max_rows, engine=None: (list(columns), list(rows)),
    )


class TestCreateConnector:
    def test_an_admin_can_create_a_power_bi_connector(self, client):
        headers = _account(client, "admin1@tenant-a.example.com", "admin")
        body = _create_connector(client, headers)
        assert body["target_type"] == "power_bi"
        assert body["status"] == "never_refreshed"
        assert body["tenant_id"] == "tenant-a"

    def test_a_non_admin_cannot_create_a_connector(self, client):
        headers = _account(client, "analyst1@tenant-a.example.com", "analyst")
        response = client.post(
            "/bi/connectors",
            json={
                "target_type": "power_bi",
                "display_name": "x",
                "sql_text": "SELECT 1",
                "push_config": {"dataset_id": "d", "table_name": "t"},
            },
            headers=headers,
        )
        assert response.status_code == 403

    def test_invalid_sql_is_rejected_at_creation(self, client):
        headers = _account(client, "admin2@tenant-a.example.com", "admin")
        response = client.post(
            "/bi/connectors",
            json={
                "target_type": "excel",
                "display_name": "Bad",
                "sql_text": "DROP TABLE sales",
                "push_config": {},
            },
            headers=headers,
        )
        assert response.status_code == 422

    def test_power_bi_connector_missing_push_config_keys_is_rejected(self, client):
        headers = _account(client, "admin3@tenant-a.example.com", "admin")
        response = client.post(
            "/bi/connectors",
            json={
                "target_type": "power_bi",
                "display_name": "Bad",
                "sql_text": "SELECT 1",
                "push_config": {},
            },
            headers=headers,
        )
        assert response.status_code == 422
        assert "dataset_id" in response.text or "table_name" in response.text

    def test_looker_connector_missing_webhook_url_is_rejected(self, client):
        headers = _account(client, "admin4@tenant-a.example.com", "admin")
        response = client.post(
            "/bi/connectors",
            json={
                "target_type": "looker",
                "display_name": "Bad",
                "sql_text": "SELECT 1",
                "push_config": {},
            },
            headers=headers,
        )
        assert response.status_code == 422

    def test_a_tableau_connector_needs_no_push_config(self, client):
        headers = _account(client, "admin5@tenant-a.example.com", "admin")
        body = _create_connector(
            client, headers, target_type="tableau", display_name="Tableau Feed", push_config={}
        )
        assert body["target_type"] == "tableau"

    def test_a_published_governed_metric_can_be_linked(self, client):
        import identity.repositories.semantic_catalog as catalog_repo

        session = identity_db_mod.get_identity_session(_SETTINGS)
        try:
            entry = catalog_repo.create_entry(
                session,
                tenant_id="tenant-a",
                database_id="default",
                concept_type="metric",
                concept_key="net_revenue",
                business_name="Net Revenue",
                technical_name=None,
                description="Gross sales minus refunds.",
                grain=None,
                keys=[],
                relationships=[],
                domain=None,
                synonyms=[],
                business_rules=[],
                examples=[],
                evidence=[],
                confidence=1.0,
                owner="finance",
                created_by_user_id=None,
                approved_expression="SUM(sales.amount)",
                source_tables=["sales"],
            )
            reviewed = catalog_repo.mark_reviewed(session, entry, reviewed_by_user_id=entry.id)
            catalog_repo.publish_entry(session, reviewed, published_by_user_id=entry.id)
            session.commit()
        finally:
            session.close()

        headers = _account(client, "admin30@tenant-a.example.com", "admin")
        body = _create_connector(
            client,
            headers,
            target_type="excel",
            display_name="Revenue Feed",
            sql_text="SELECT region, amount FROM sales",
            push_config={},
            governed_metric_concept_key="net_revenue",
        )
        assert body["governed_metric_concept_key"] == "net_revenue"
        assert body["governed_metric_version"] == 1

    def test_an_unknown_governed_metric_is_rejected(self, client):
        headers = _account(client, "admin6@tenant-a.example.com", "admin")
        response = client.post(
            "/bi/connectors",
            json={
                "target_type": "excel",
                "display_name": "x",
                "sql_text": "SELECT 1",
                "push_config": {},
                "governed_metric_concept_key": "does_not_exist",
            },
            headers=headers,
        )
        assert response.status_code == 422

    def test_bi_integration_disabled_is_a_404(self, client, monkeypatch):
        headers = _account(client, "admin7@tenant-a.example.com", "admin")
        disabled_settings = _SETTINGS.model_copy(update={"enable_bi_integration": False})
        monkeypatch.setattr(bi_integration_mod, "get_settings", lambda: disabled_settings)
        response = client.post(
            "/bi/connectors",
            json={
                "target_type": "excel",
                "display_name": "x",
                "sql_text": "SELECT 1",
                "push_config": {},
            },
            headers=headers,
        )
        assert response.status_code == 404


class TestListGetDelete:
    def test_list_is_tenant_scoped(self, client):
        headers_a = _account(client, "admin8@tenant-a.example.com", "admin")
        headers_b = _account(client, "admin9@tenant-b.example.com", "admin", tenant_id="tenant-b")
        _create_connector(client, headers_a)
        _create_connector(client, headers_b)
        assert len(client.get("/bi/connectors", headers=headers_a).json()) == 1
        assert len(client.get("/bi/connectors", headers=headers_b).json()) == 1

    def test_get_cross_tenant_connector_is_a_404(self, client):
        headers_a = _account(client, "admin10@tenant-a.example.com", "admin")
        headers_b = _account(client, "admin11@tenant-b.example.com", "admin", tenant_id="tenant-b")
        connector = _create_connector(client, headers_a)
        response = client.get(f"/bi/connectors/{connector['id']}", headers=headers_b)
        assert response.status_code == 404

    def test_get_nonexistent_connector_is_also_a_404(self, client):
        headers = _account(client, "admin12@tenant-a.example.com", "admin")
        response = client.get(
            "/bi/connectors/00000000-0000-0000-0000-000000000000", headers=headers
        )
        assert response.status_code == 404

    def test_delete_removes_the_connector(self, client):
        headers = _account(client, "admin13@tenant-a.example.com", "admin")
        connector = _create_connector(client, headers)
        delete_response = client.delete(f"/bi/connectors/{connector['id']}", headers=headers)
        assert delete_response.status_code == 204
        get_response = client.get(f"/bi/connectors/{connector['id']}", headers=headers)
        assert get_response.status_code == 404


class TestRefreshSuccess:
    def test_refreshing_a_power_bi_connector_pushes_rows(self, client, monkeypatch):
        headers = _account(client, "admin14@tenant-a.example.com", "admin")
        connector = _create_connector(client, headers)
        _mock_execute_success(monkeypatch)
        monkeypatch.setattr(
            bi_integration_mod,
            "push_power_bi",
            lambda settings, dataset_id, table_name, columns, rows: PushResult(
                success=True, detail="ok", row_count=len(rows)
            ),
        )
        response = client.post(f"/bi/connectors/{connector['id']}/refresh", headers=headers)
        assert response.status_code == 200, response.text
        body = response.json()
        assert body["status"] == "healthy"
        assert body["last_refresh_status"] == "success"
        assert body["last_row_count"] == 1

    def test_refreshing_a_pull_connector_populates_the_export_cache(self, client, monkeypatch):
        headers = _account(client, "admin15@tenant-a.example.com", "admin")
        connector = _create_connector(
            client, headers, target_type="excel", display_name="Excel Feed", push_config={}
        )
        _mock_execute_success(monkeypatch, columns=("x",), rows=((1,), (2,)))
        refresh_response = client.post(f"/bi/connectors/{connector['id']}/refresh", headers=headers)
        assert refresh_response.status_code == 200
        assert refresh_response.json()["status"] == "healthy"

        export_response = client.get(
            f"/bi/connectors/{connector['id']}/export?format=csv", headers=headers
        )
        assert export_response.status_code == 200
        # csv.writer's own default dialect uses \r\n line terminators
        # (RFC 4180) -- this is the correct, standard CSV shape, not a bug.
        assert export_response.text == "x\r\n1\r\n2\r\n"

    def test_export_before_any_refresh_is_a_409(self, client):
        headers = _account(client, "admin16@tenant-a.example.com", "admin")
        connector = _create_connector(
            client, headers, target_type="tableau", display_name="Feed", push_config={}
        )
        response = client.get(
            f"/bi/connectors/{connector['id']}/export?format=csv", headers=headers
        )
        assert response.status_code == 409

    def test_export_of_a_push_connector_is_a_409(self, client, monkeypatch):
        headers = _account(client, "admin17@tenant-a.example.com", "admin")
        connector = _create_connector(client, headers)
        response = client.get(
            f"/bi/connectors/{connector['id']}/export?format=csv", headers=headers
        )
        assert response.status_code == 409

    def test_export_xlsx_format(self, client, monkeypatch):
        import io

        from openpyxl import load_workbook

        headers = _account(client, "admin18@tenant-a.example.com", "admin")
        connector = _create_connector(
            client, headers, target_type="excel", display_name="Feed", push_config={}
        )
        _mock_execute_success(monkeypatch, columns=("x", "y"), rows=((1, 2),))
        client.post(f"/bi/connectors/{connector['id']}/refresh", headers=headers)
        response = client.get(
            f"/bi/connectors/{connector['id']}/export?format=xlsx", headers=headers
        )
        assert response.status_code == 200
        workbook = load_workbook(io.BytesIO(response.content))
        rows = list(workbook.active.iter_rows(values_only=True))
        assert rows[0] == ("x", "y")
        assert rows[1] == (1, 2)


class TestRefreshFailures:
    """The mission's own explicit 'Test authorization and refresh
    failures' requirement."""

    def test_a_sql_execution_error_is_reported_as_a_200_with_error_status(
        self, client, monkeypatch
    ):
        headers = _account(client, "admin19@tenant-a.example.com", "admin")
        connector = _create_connector(client, headers)

        def _raise(sql, timeout, max_rows, engine=None):
            raise SQLAlchemyError("connection refused")

        monkeypatch.setattr(bi_integration_mod, "execute_readonly_sql", _raise)
        response = client.post(f"/bi/connectors/{connector['id']}/refresh", headers=headers)
        assert response.status_code == 200
        body = response.json()
        assert body["status"] == "error"
        assert body["last_refresh_status"] == "failed"
        assert "connection refused" in body["last_error"]

    def test_a_push_failure_is_reported_as_a_200_with_error_status(self, client, monkeypatch):
        headers = _account(client, "admin20@tenant-a.example.com", "admin")
        connector = _create_connector(client, headers)
        _mock_execute_success(monkeypatch)
        from bi_integration.push import BiPushNotConfiguredError

        def _raise(*args, **kwargs):
            raise BiPushNotConfiguredError("Power BI is not configured.")

        monkeypatch.setattr(bi_integration_mod, "push_power_bi", _raise)
        response = client.post(f"/bi/connectors/{connector['id']}/refresh", headers=headers)
        assert response.status_code == 200
        body = response.json()
        assert body["status"] == "error"
        assert "not configured" in body["last_error"]

    def test_a_webhook_ssrf_rejection_is_reported_as_a_200_with_error_status(
        self, client, monkeypatch
    ):
        headers = _account(client, "admin21@tenant-a.example.com", "admin")
        connector = _create_connector(
            client,
            headers,
            target_type="looker",
            display_name="Looker",
            push_config={"webhook_url": "https://internal.example/webhook"},
        )
        _mock_execute_success(monkeypatch)
        from bi_integration.outbound_safety import OutboundUrlBlockedError

        def _raise(*args, **kwargs):
            raise OutboundUrlBlockedError("Refusing to call a private address")

        monkeypatch.setattr(bi_integration_mod, "push_looker_webhook", _raise)
        response = client.post(f"/bi/connectors/{connector['id']}/refresh", headers=headers)
        assert response.status_code == 200
        body = response.json()
        assert body["status"] == "error"
        assert "private address" in body["last_error"]

    def test_a_removed_database_fails_the_refresh_gracefully(self, client, monkeypatch):
        headers = _account(client, "admin22@tenant-a.example.com", "admin")
        connector = _create_connector(client, headers)
        # Simulates this tenant's own configured database having been
        # removed/renamed since the connector was created -- the refresh
        # must degrade to a recorded error, never a 500.
        monkeypatch.setattr(Settings, "databases_for_tenant", lambda self, tenant_id: ())
        response = client.post(f"/bi/connectors/{connector['id']}/refresh", headers=headers)
        assert response.status_code == 200
        body = response.json()
        assert body["status"] == "error"
        assert "no longer configured" in body["last_error"]

    def test_refresh_requires_bi_integration_manage(self, client):
        admin_headers = _account(client, "admin23@tenant-a.example.com", "admin")
        analyst_headers = _account(client, "analyst2@tenant-a.example.com", "analyst")
        connector = _create_connector(client, admin_headers)
        response = client.post(f"/bi/connectors/{connector['id']}/refresh", headers=analyst_headers)
        assert response.status_code == 403

    def test_refresh_of_a_cross_tenant_connector_is_a_404(self, client):
        headers_a = _account(client, "admin24@tenant-a.example.com", "admin")
        headers_b = _account(client, "admin25@tenant-b.example.com", "admin", tenant_id="tenant-b")
        connector = _create_connector(client, headers_a)
        response = client.post(f"/bi/connectors/{connector['id']}/refresh", headers=headers_b)
        assert response.status_code == 404


class TestUnauthorizedAccess:
    def test_every_route_requires_authentication(self, client):
        assert client.get("/bi/connectors").status_code in (401, 403)
        assert client.post(
            "/bi/connectors",
            json={
                "target_type": "excel",
                "display_name": "x",
                "sql_text": "SELECT 1",
                "push_config": {},
            },
        ).status_code in (401, 403)

    def test_a_viewer_cannot_delete_a_connector(self, client):
        admin_headers = _account(client, "admin26@tenant-a.example.com", "admin")
        viewer_headers = _account(client, "viewer1@tenant-a.example.com", "viewer")
        connector = _create_connector(client, admin_headers)
        response = client.delete(f"/bi/connectors/{connector['id']}", headers=viewer_headers)
        assert response.status_code == 403
