"""Unit tests for the pure `bi_integration/` modules -- Prompt 48
(`48_ENTERPRISE_BI_INTEGRATIONS_CONTRACT.md`). No FastAPI, no identity
database -- see `tests/test_api_bi_integration.py` for the HTTP-level
coverage and `tests/test_identity_bi_integration_repository.py` for the
persistence layer.
"""

from __future__ import annotations

import io
from datetime import date
from decimal import Decimal
from unittest.mock import patch

import pytest
from identity.rbac import Permission
from openpyxl import load_workbook
from semantic.catalog import CatalogEntrySnapshot, CatalogStatus

from bi_integration.metric_binding import referenced_tables, validate_governed_metric_binding
from bi_integration.models import PULL_TARGETS, PUSH_TARGETS, BiTargetType
from bi_integration.outbound_safety import OutboundUrlBlockedError, assert_safe_outbound_url
from bi_integration.policy import BiAction, authorize_bi_action
from bi_integration.push import BiPushNotConfiguredError, push_looker_webhook, push_power_bi
from bi_integration.rendering import render_csv, render_json, render_xlsx
from config.settings import Settings


class TestModels:
    def test_push_and_pull_targets_partition_cleanly(self):
        assert set(BiTargetType) == PUSH_TARGETS | PULL_TARGETS
        assert set() == PUSH_TARGETS & PULL_TARGETS

    def test_push_targets_are_power_bi_and_looker(self):
        assert {BiTargetType.POWER_BI, BiTargetType.LOOKER} == PUSH_TARGETS


class TestMetricBinding:
    def test_referenced_tables_handles_a_join(self):
        tables = referenced_tables("SELECT * FROM sales s JOIN customers c ON s.cid = c.id")
        assert tables == frozenset({"sales", "customers"})

    def test_referenced_tables_returns_empty_on_unparseable_sql(self):
        assert referenced_tables("not valid sql (((") == frozenset()

    def _snapshot(self, **overrides) -> CatalogEntrySnapshot:
        base = dict(
            id="x",
            tenant_id="t1",
            database_id="default",
            concept_type="metric",
            concept_key="net_revenue",
            business_name="Net Revenue",
            status=CatalogStatus.PUBLISHED,
            source_tables=("sales",),
            version=1,
        )
        base.update(overrides)
        return CatalogEntrySnapshot(**base)

    def test_none_snapshot_is_always_accepted(self):
        assert validate_governed_metric_binding(None, "SELECT 1") == []

    def test_published_with_overlapping_tables_is_accepted(self):
        snapshot = self._snapshot()
        assert validate_governed_metric_binding(snapshot, "SELECT SUM(amount) FROM sales") == []

    def test_published_with_no_table_overlap_is_rejected(self):
        snapshot = self._snapshot()
        violations = validate_governed_metric_binding(snapshot, "SELECT SUM(amount) FROM orders")
        assert len(violations) == 1
        assert "sales" in violations[0]

    def test_draft_status_is_rejected_even_with_table_overlap(self):
        snapshot = self._snapshot(status=CatalogStatus.DRAFT)
        violations = validate_governed_metric_binding(snapshot, "SELECT SUM(amount) FROM sales")
        assert any("not published" in v for v in violations)

    def test_a_snapshot_with_no_declared_source_tables_skips_the_overlap_check(self):
        snapshot = self._snapshot(source_tables=())
        assert validate_governed_metric_binding(snapshot, "SELECT SUM(amount) FROM orders") == []


class TestPolicy:
    def test_denies_when_actor_has_no_resolved_tenant(self):
        decision = authorize_bi_action(
            actor_permissions=frozenset({Permission.BI_INTEGRATION_MANAGE.value}),
            actor_tenant_id=None,
            resource_tenant_id=None,
            action=BiAction.LIST,
        )
        assert not decision.allowed

    def test_cross_tenant_resource_is_denied_before_any_permission_check(self):
        decision = authorize_bi_action(
            actor_permissions=frozenset({Permission.BI_INTEGRATION_MANAGE.value}),
            actor_tenant_id="tenant-a",
            resource_tenant_id="tenant-b",
            action=BiAction.GET,
        )
        assert not decision.allowed
        assert "different tenant" in decision.reason

    def test_no_permission_is_denied(self):
        decision = authorize_bi_action(
            actor_permissions=frozenset(),
            actor_tenant_id="tenant-a",
            resource_tenant_id=None,
            action=BiAction.CREATE,
        )
        assert not decision.allowed

    def test_bi_integration_manage_is_sufficient_for_every_action(self):
        for action in BiAction:
            decision = authorize_bi_action(
                actor_permissions=frozenset({Permission.BI_INTEGRATION_MANAGE.value}),
                actor_tenant_id="tenant-a",
                resource_tenant_id=None,
                action=action,
            )
            assert decision.allowed, action

    def test_an_unrecognized_action_is_denied(self):
        decision = authorize_bi_action(
            actor_permissions=frozenset({Permission.BI_INTEGRATION_MANAGE.value}),
            actor_tenant_id="tenant-a",
            resource_tenant_id=None,
            action="not_a_real_action",  # type: ignore[arg-type]
        )
        assert not decision.allowed
        assert decision.reason == "unrecognized action"


class TestRendering:
    _COLUMNS = ["region", "revenue", "as_of"]
    _ROWS = [
        ("East", Decimal("1234.5"), date(2026, 1, 1)),
        ("West", Decimal("999.9"), date(2026, 1, 1)),
    ]

    def test_render_csv_round_trips(self):
        text = render_csv(self._COLUMNS, self._ROWS).decode()
        assert "region,revenue,as_of" in text
        assert "East,1234.5,2026-01-01" in text

    def test_render_json_coerces_decimal_and_date(self):
        import json

        payload = json.loads(render_json(self._COLUMNS, self._ROWS).decode())
        assert payload[0] == {"region": "East", "revenue": 1234.5, "as_of": "2026-01-01"}

    def test_render_xlsx_round_trips(self):
        data = render_xlsx(self._COLUMNS, self._ROWS, "Revenue")
        workbook = load_workbook(io.BytesIO(data))
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        assert rows[0] == ("region", "revenue", "as_of")
        assert rows[1] == ("East", 1234.5, "2026-01-01")

    def test_render_xlsx_truncates_a_too_long_sheet_title(self):
        data = render_xlsx(self._COLUMNS, self._ROWS, "x" * 50)
        workbook = load_workbook(io.BytesIO(data))
        assert len(workbook.active.title) <= 31

    def test_render_xlsx_row_limit_is_enforced(self, monkeypatch):
        import bi_integration.rendering as rendering_mod

        monkeypatch.setattr(rendering_mod, "_XLSX_MAX_ROWS", 1)
        with pytest.raises(ValueError, match="exceeds"):
            rendering_mod.render_xlsx(["x"], [(1,), (2,)], "t")


class TestOutboundSafety:
    def test_non_https_is_rejected(self):
        with pytest.raises(OutboundUrlBlockedError, match="non-HTTPS"):
            assert_safe_outbound_url("http://example.com/webhook", context="test")

    def test_no_hostname_is_rejected(self):
        with pytest.raises(OutboundUrlBlockedError, match="no hostname"):
            assert_safe_outbound_url("https:///webhook", context="test")

    def test_private_address_is_rejected(self):
        with (
            patch("socket.getaddrinfo", return_value=[(2, 1, 6, "", ("10.0.0.5", 0))]),
            pytest.raises(OutboundUrlBlockedError, match="private/internal"),
        ):
            assert_safe_outbound_url("https://malicious.example/webhook", context="test")

    def test_public_address_is_allowed(self):
        with patch("socket.getaddrinfo", return_value=[(2, 1, 6, "", ("93.184.216.34", 0))]):
            assert_safe_outbound_url("https://public.example/webhook", context="test")  # no raise

    def test_rejection_is_audit_logged(self):
        with (
            patch("bi_integration.outbound_safety.log_security_event") as mock_log,
            pytest.raises(OutboundUrlBlockedError),
        ):
            assert_safe_outbound_url("http://example.com", context="my_context")
        mock_log.assert_called_once()
        args, kwargs = mock_log.call_args
        assert args[0] == "ssrf_blocked"
        assert kwargs["context"] == "my_context"

    def test_unresolvable_host_is_not_logged_as_ssrf_blocked(self):
        import socket as socket_module

        with (
            patch("socket.getaddrinfo", side_effect=socket_module.gaierror("nope")),
            patch("bi_integration.outbound_safety.log_security_event") as mock_log,
            pytest.raises(OutboundUrlBlockedError),
        ):
            assert_safe_outbound_url("https://does-not-exist.invalid/webhook", context="t")
        mock_log.assert_not_called()


_PUBLIC_ADDRINFO = [(2, 1, 6, "", ("93.184.216.34", 0))]


class TestPush:
    def test_power_bi_push_raises_when_not_configured(self):
        settings = Settings(_env_file=None)
        with pytest.raises(BiPushNotConfiguredError):
            push_power_bi(settings, "ds1", "tbl1", ["x"], [(1,)])

    def test_power_bi_push_posts_the_documented_shape(self):
        from unittest.mock import MagicMock

        settings = Settings(
            _env_file=None,
            power_bi_push_base_url="https://api.powerbi.com/v1.0/myorg",
            power_bi_push_api_key="fakekey",
        )
        mock_response = MagicMock()
        mock_response.raise_for_status.return_value = None
        with (
            patch("socket.getaddrinfo", return_value=_PUBLIC_ADDRINFO),
            patch("httpx.post", return_value=mock_response) as mock_post,
        ):
            result = push_power_bi(settings, "ds1", "tbl1", ["x", "y"], [(1, 2)])
        assert result.success
        assert result.row_count == 1
        url, kwargs = mock_post.call_args.args[0], mock_post.call_args.kwargs
        assert url == "https://api.powerbi.com/v1.0/myorg/datasets/ds1/tables/tbl1/rows"
        assert kwargs["json"] == {"rows": [{"x": 1, "y": 2}]}
        assert kwargs["headers"]["Authorization"] == "Bearer fakekey"

    def test_power_bi_push_rejects_a_private_base_url(self):
        settings = Settings(
            _env_file=None,
            power_bi_push_base_url="https://internal.example",
            power_bi_push_api_key="fakekey",
        )
        with (
            patch("socket.getaddrinfo", return_value=[(2, 1, 6, "", ("10.0.0.5", 0))]),
            pytest.raises(OutboundUrlBlockedError),
        ):
            push_power_bi(settings, "ds1", "tbl1", ["x"], [(1,)])

    def test_looker_webhook_push_posts_the_documented_shape(self):
        from unittest.mock import MagicMock

        settings = Settings(_env_file=None)
        mock_response = MagicMock()
        mock_response.raise_for_status.return_value = None
        with (
            patch("socket.getaddrinfo", return_value=_PUBLIC_ADDRINFO),
            patch("httpx.post", return_value=mock_response) as mock_post,
        ):
            result = push_looker_webhook(
                settings, "https://company.cloud.looker.com/webhook/abc", ["x"], [(1,)]
            )
        assert result.success
        kwargs = mock_post.call_args.kwargs
        assert kwargs["json"] == {"fields": ["x"], "data": [{"x": 1}]}

    def test_looker_webhook_push_rejects_a_non_https_url(self):
        settings = Settings(_env_file=None)
        with pytest.raises(OutboundUrlBlockedError):
            push_looker_webhook(settings, "http://company.looker.com/webhook", ["x"], [(1,)])
